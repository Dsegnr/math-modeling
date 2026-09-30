# -*- coding: utf-8 -*-
"""在现有 docx 初稿上原地扩充公式（只新增公式段落 + 重排编号，不动其他内容）。"""
import re
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import parse_xml

M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
MATH = "{%s}oMath" % M_NS


def mr(t):
    return '<m:r><m:t xml:space="preserve">%s</m:t></m:r>' % (
        t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


def make_eq_para(doc, inner):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.tab_stops.add_tab_stop(Cm(16.0),
                                              WD_TAB_ALIGNMENT.RIGHT)
    omath = parse_xml('<m:oMath xmlns:m="%s" xmlns:w="%s">%s</m:oMath>'
                      % (M_NS, W_NS, inner))
    p._p.append(omath)
    r = p.add_run("\t（公式0）")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    return p


def find_para(doc, substr):
    for p in doc.paragraphs:
        if substr in p.text:
            return p
    return None


def insert_after(anchor, new_p):
    anchor._p.addnext(new_p._p)
    return new_p


def renumber(doc):
    n = 0
    for p in doc.paragraphs:
        if p._p.find(MATH) is not None:
            n += 1
            for r in p.runs:
                if re.search(r"（公式\s*\d+）", r.text):
                    r.text = re.sub(r"（公式\s*\d+）", "（公式%d）" % n, r.text)
    return n


def expand_ch58(path):
    doc = docx.Document(path)
    missing = []

    def anchor(sub):
        p = find_para(doc, sub)
        if p is None:
            missing.append(sub)
        return p

    # B：混合体系阈值（A-B/B-B/B-面）——公式2 之后
    a = anchor("（公式2）")
    if a is not None:
        insert_after(a, make_eq_para(
            doc, mr("d_AB = R_A+R_B+δ = 231.8 nm，d_BB = 2R_B+δ = 401.8 nm，"
                    "d_BF = R_B+δ = 201.8 nm")))
    # C：线段-线段最短距离形式化
    a = anchor("线段-线段最短距离采用三维参数化最近点对求解")
    if a is not None:
        insert_after(a, make_eq_para(
            doc, mr("d(s_i, s_j) = min ‖s_i(u) − s_j(v)‖，u, v ∈ [0, 1]")))
    # D：图模型定义
    a = anchor("边集合 E 由接触判定生成")
    if a is not None:
        insert_after(a, make_eq_para(
            doc, mr("G = (V, E)，V = {s_1, …, s_N, L, R}，"
                    "E = {(u, v) | d(u, v) ≤ TH}")))
    # E：连通判据
    a = anchor("存在一条从 L 出发、经介质段节点、到达 R 的路径")
    if a is not None:
        insert_after(a, make_eq_para(
            doc, mr("导通 ⇔ find(L) = find(R)（L 与 R 同根）")))
    # A：体积公式——公式3 之后
    a = anchor("（公式3）")
    if a is not None:
        insert_after(a, make_eq_para(
            doc, mr("V_C = a³ = 10¹² nm³，V_A = πR_A²H_A ≈ 1.414×10⁷ nm³，"
                    "V_B = (4/3)πR_B³ ≈ 3.35×10⁷ nm³")))
    # F：二项方差——公式5 之后
    a = anchor("（公式5）")
    if a is not None:
        insert_after(a, make_eq_para(
            doc, mr("Var(P̂) ≈ P(1−P)/M")))
    # G/H：加权最小二乘与拟合反解——公式8 之后（先后插入，保持顺序）
    a = anchor("（公式8）")
    if a is not None:
        g = insert_after(a, make_eq_para(
            doc, mr("min Σ w_i [P̂_i − P(φ_i; θ)]²")))
        insert_after(g, make_eq_para(doc, mr("P(φ*) = 0.90")))
    # I：二分更新——公式10 之后
    a = anchor("（公式10）")
    if a is not None:
        insert_after(a, make_eq_para(
            doc, mr("N_mid = ⌊(N_lo + N_hi)/2⌋，按 CI 下界判可行并收缩区间")))
    # J/M/N：等效体积、纯B理论成本、占优条件——公式13 之后
    a = anchor("（公式13）")
    if a is not None:
        j = insert_after(a, make_eq_para(
            doc, mr("w = φ*_A / φ*_B ≈ 0.022")))
        m = insert_after(j, make_eq_para(
            doc, mr("C_B = 1000 · c_B · (φ*_A / w)")))
        insert_after(m, make_eq_para(
            doc, mr("纯 A 占优条件：c_A·φ*_A < c_B·φ*_A/w ⇔ w < c_B/c_A ≈ 0.0476")))
    # K/L：罚函数与面接触瓶颈——公式14 之后
    a = anchor("（公式14）")
    if a is not None:
        k = insert_after(a, make_eq_para(
            doc, mr("min C(φA,φB) + λ·max(0, 0.90 − P_surf(φA,φB))")))
        insert_after(k, make_eq_para(
            doc, mr("t_B,eff = (R_B + δ) − R_B = δ = 1.8 nm")))

    total = renumber(doc)
    doc.save(path)
    return total, missing


def expand_ch9(path):
    doc = docx.Document(path)
    missing = []

    def anchor(sub):
        p = find_para(doc, sub)
        if p is None:
            missing.append(sub)
        return p

    # Q：口径差异——9.2 正文后
    a = anchor("仅 C1 给出")
    if a is not None:
        insert_after(a, make_eq_para(
            doc, mr("ΔP(φ) = P_A1(φ) − P_C1(φ)（0.5% 处约 0.94）")))
    # P：最小样本量——公式16 之后
    a = anchor("（公式16）")
    if a is not None:
        insert_after(a, make_eq_para(
            doc, mr("M ≥ (z/ε)² · P̂(1−P̂)，ε 为允许的区间半宽")))
    # O：灵敏度指数——公式17 之后
    a = anchor("（公式17）")
    if a is not None:
        insert_after(a, make_eq_para(
            doc, mr("S = (Δφ*/φ*) / (Δf/f)")))
    # R：成本比——公式18 之后
    a = anchor("（公式18）")
    if a is not None:
        insert_after(a, make_eq_para(
            doc, mr("ρ = c_B·φ*_B / (c_A·φ*_A) ≈ 2.15")))

    total = renumber(doc)
    doc.save(path)
    return total, missing


if __name__ == "__main__":
    p58 = (r"D:\数模竞赛\华数杯2026\A题项目\华数杯2026_A题\写作素材"
           r"\五至八章_模型建立与求解_初稿.docx")
    p9 = (r"D:\数模竞赛\华数杯2026\A题项目\华数杯2026_A题\写作素材"
          r"\第九章_模型的分析与检验_初稿.docx")
    n58, m58 = expand_ch58(p58)
    n9, m9 = expand_ch9(p9)
    print("ch5-8 formulas:", n58, "missing anchors:", m58)
    print("ch9 formulas:", n9, "missing anchors:", m9)
