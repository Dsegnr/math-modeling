# -*- coding: utf-8 -*-
"""华数杯 A 题 hsb(8).pdf -> docx 完整转换稿的定点精修。

修复清单：
1) 摘要压缩为 6 段（控制在一页内）；
2) 删除全部纯数字页码碎片（pdf2docx 把页脚当正文段）；
3) 修复 9.4（问题四检验）复制 9.3 的错误；
4) 10.1 标题“模型的缺点”->“模型的优点”（内容本就是优点）；
5) 表8/表10 标题重复编号；
6) 图15 图注误标“问题三”->“问题二”；
7) 9.3 正文“图30 和表16”错误引用 -> “图16 和表12”；
8) 5.5 重复口径句删除；
9) 9.2 多随机种子段落断行拼接（第23页跨页）；
10) 图片宽度 >13cm 等比收缩；
11) 图15(a)(c) 图注细节修正。
"""
import copy
import re
import docx
from docx.shared import Cm, Emu

SRC = r"D:\数模竞赛\华数杯2026\A题项目\华数杯2026_A题\写作素材\hsb8_full_convert.docx"
DST = r"D:\数模竞赛\华数杯2026\A题项目\华数杯2026_A题\写作素材\hsb8_final.docx"

d = docx.Document(SRC)
paras = d.paragraphs


def set_para_text(par, text):
    """保留首 run 格式，整体替换段落文本。"""
    if not par.runs:
        par.add_run(text)
        return
    par.runs[0].text = text
    for r in par.runs[1:]:
        r.text = ""


def find_par(prefix):
    for p in paras:
        if p.text.strip().startswith(prefix):
            return p
    return None


def remove_par(par):
    parent = par._p.getparent()
    if parent is not None:
        parent.remove(par._p)


log = []

# ---------- 1) 摘要压缩 ----------
abstract_new = [
    "可靠且可复现的导通性评估对导电填充材料的设计与成本优化至关重要。本文针对微构体中"
    "随机填充圆柱介质A与球形介质B的导通判定、导通概率估计、临界体积分数反求与双介质"
    "最低成本设计问题，以三维最短距离与并查集构建四问共用的确定性判定内核，结合拒绝采样"
    "蒙特卡洛、渗透模型加权拟合、基于Wilson区间下界的概率二分、等效体积线性规划与Logistic"
    "代理优化求解，并将全部可行性结论统一置于“95%置信区间下界不低于0.90”的判据之下。",
    "针对问题一，将附件各行介质A抽象为轴线段，以左右带电面为源汇节点构造接触图，按题面"
    "表面间距判据换算轴线距离阈值，用向量化三维线段最近点算法批量生成接触边，以并查集"
    "判定连通并回溯跨接路径、定位断口。结果表明，主口径下组1因约279nm的关键断口不导通，"
    "组2、组3均存在完整通路；同根合并口径三组全导通，作为口径敏感性对照列入第九章。",
    "针对问题二，在介质完全位于盒内的C1口径下以拒绝采样生成随机构型，复用问题一内核做"
    "蒙特卡洛估计（每点3000次，多进程并行），以Wilson得分区间、收敛曲线与多种子复核刻画"
    "统计可靠性。结果表明，导通概率呈典型渗流相变，由0.5%时的0.060经1.0%时的0.730升至"
    "1.5%时的0.949；卷回口径（A1/B1）在相同区间恒接近1，与题目非平凡概率设定不符，故"
    "主口径取C1。",
    "针对问题三，以Wilson区间下界不低于0.90为统一可行标准：先拟合Sigmoid与幂律截断渗透"
    "模型（加权最小二乘、AIC选优）并解析反解，再以整数介质数量为变量做概率二分，随后在"
    "两解邻域高样本细化并多种子终验。最终定稿最低体积分数φ*=1.359%（N=961），终验"
    "P̂=0.911，95%CI=[0.905,0.917]，拟合质量另以残差诊断核对。",
    "针对问题四，构建“理论层＋工程层＋验证层”三层优化框架：理论层由纯A、纯B临界体积"
    "分数标定等效体积权重并建立线性规划求成本下界；工程层拟合Logistic代理曲面并以罚函数"
    "SLSQP搜索混合候选；验证层对全部候选做真实蒙特卡洛复核，低成本混合候选因此被否决。"
    "最终最优方案为纯A，φ*=1.359%，成本14.27元，不足纯B方案的一半；机理上球形介质存在"
    "“面接触瓶颈”，球体有效贴面厚度仅约1.8nm，纯B在60%体积分数下仍不满足判据。",
    "最后，对口径选择、导通阈值、样本量收敛、多种子复现及单价与等效权重扰动进行系统敏感性"
    "分析，全部结论稳健且可复现（固定随机流、逐样本落盘、断点续跑）；方法可推广至复合材料"
    "导电网络、多孔介质渗流及电池电极导电剂网络等随机微结构体系。",
]
start = find_par("可靠且可复现的导通性评估")
kw = find_par("关键词")
if start is not None and kw is not None:
    pars = []
    on = False
    for p in paras:
        if p is start:
            on = True
        if on:
            pars.append(p)
        if p is kw:
            break
    for i, p in enumerate(pars):
        if i < len(abstract_new):
            set_para_text(p, abstract_new[i])
        else:
            remove_par(p)
    log.append(f"摘要：原 {len(pars)} 段 -> {len(abstract_new)} 段")
else:
    log.append("摘要：未找到锚点！")


# ---------- 2) 删除纯数字页码碎片 ----------
removed_pages = 0
for p in list(paras):
    t = p.text.strip()
    if re.fullmatch(r"\d{1,3}", t):
        remove_par(p)
        removed_pages += 1
log.append(f"页码碎片：删除 {removed_pages} 段")


# ---------- 3) 9.4 修复（问题四检验） ----------
h94 = None
for p in paras:
    if p.text.strip().startswith("9.4"):
        h94 = p
        break
if h94 is not None:
    # 保留标题 run（0..5），替换正文 run
    runs = h94.runs
    if len(runs) >= 6:
        # runs[0]='9.4' runs[2..3]='\t\t问题四模型的检验' runs[4]='\n' runs[5]='\t'
        runs[3].text = "问题四模型的检验"
        new_body = (
            "问题四重点检验介质单价、等效体积权重扰动及纯B外推不确定性。对介质A、B单价分别"
            "做±10%扰动并重解理论层线性规划，全部组合下LP极点均退化为纯A（成本12.84~15.70元），"
            "最优方案不随单价扰动改变；将等效体积权重w在±20%内扰动，LP最优恒为纯A（成本14.27元），"
            "说明成本结论对w不敏感。纯B扫描上限为60%，φ*_B≈61.4%为外推插值；50%、55%、60%三点"
            "真实MC结果的CI下界分别为0.721、0.817、0.856，均低于0.90，外推处P̂=0.902但CI下界"
            "0.882仍不达标，因此“纯B不达标”结论保守稳健。相关结果见图17。"
        )
        runs[6].text = new_body
        for r in runs[7:]:
            r.text = ""
        # 删除 9.4 与 9.5 之间的复制内容段（含页码/空段已在上一步删掉）
        content = []
        on = False
        for p in paras:
            if p is h94:
                on = True
                continue
            if on:
                if p.text.strip().startswith("9.5"):
                    break
                content.append(p)
        removed = 0
        for p in content:
            t = p.text.strip()
            if t.startswith("接触阈值缩放因子") or t.startswith("利用阈值敏感性曲线"):
                remove_par(p)
                removed += 1
        log.append(f"9.4：正文已替换，删除复制段 {removed} 段")
else:
    log.append("9.4：未找到！")


# ---------- 4) 10.1 标题 ----------
for p in paras:
    for r in p.runs:
        if "模型的缺点" in r.text and p.text.strip().startswith("10.1"):
            r.text = r.text.replace("模型的缺点", "模型的优点", 1)
            log.append("10.1：标题已改为“模型的优点”")
            break


# ---------- 5) 表8/表10 标题 ----------
for p in paras:
    t = p.text.strip()
    if t.startswith("表8:表8"):
        set_para_text(p, t.replace("表8:表8", "表8:", 1))
        log.append("表8：标题重复编号已修")
    if t.startswith("表10:表12"):
        set_para_text(p, t.replace("表10:表12", "表10:", 1))
        log.append("表10：标题重复编号已修")


# ---------- 6) 图15 图注 ----------
for p in paras:
    if p.text.strip().startswith("图15") and "问题三敏感性" in p.text:
        set_para_text(p, p.text.replace("问题三敏感性", "问题二敏感性", 1))
        log.append("图15：图注“问题三”已改为“问题二”")
        break


# ---------- 7) 9.3 错误引用“图30 和表16” ----------
fixed_ref = 0
for p in paras:
    if "结果见图30 和表16" in p.text or "结果见图30和表16" in p.text:
        set_para_text(p, p.text.replace("见图30 和表16", "见图16 和表12", 1).replace("见图30和表16", "见图16和表12", 1))
        fixed_ref += 1
log.append(f"引用修正：{fixed_ref} 处“图30/表16”->“图16/表12”")


# ---------- 8) 5.5 重复口径句 ----------
for p in paras:
    t = p.text.strip()
    if t.startswith("组1 的判定按附件行独立口径给出"):
        remove_par(p)
        log.append("5.5：重复口径句已删")
        break


# ---------- 9) 9.2 多随机种子断行拼接 ----------
for p in paras:
    t = p.text
    if "其次，对φ=" in t:
        # 找后续“0.9%1.0%1.2%”段并合并
        nxt = None
        idx = paras.index(p)
        for q in paras[idx + 1:]:
            if q.text.strip().startswith("0.9%1.0%1.2%"):
                nxt = q
                break
        if nxt is not None:
            tail = nxt.text.strip()
            tail = tail.replace("0.9%1.0%1.2%", "0.9%、1.0%、1.2%", 1)
            set_para_text(p, t + tail)
            remove_par(nxt)
            log.append("9.2：多随机种子断行已拼接")
        break


# ---------- 10) 图片收缩 ----------
target = Cm(13.0)
shrunk = 0
for sh in d.inline_shapes:
    if sh.width and sh.width > target:
        ratio = target / sh.width
        sh.width = target
        sh.height = Emu(int(sh.height * ratio))
        shrunk += 1
log.append(f"图片：{shrunk} 张 >13cm 已等比缩至 13cm")


# ---------- 11) 图15(a)/(c) 图注细节 ----------
for p in paras:
    t = p.text
    if "(a) 问题二样本量-精度收敛（M=500 3000" in t:
        set_para_text(p, t.replace("M=500 3000", "M=500~3000", 1))
        log.append("图15(a)：M=500~3000 修正")
    if "（c）生成口径敏感性（A1/B1 卷回vs C1 完全悬\n浮）" in t:
        set_para_text(p, "（c）生成口径敏感性（A1/B1 卷回 vs C1 完全悬浮）")
        log.append("图15(c)：图注排版修正")


d.save(DST)
print("\n".join(log))
print("SAVED:", DST)
