# -*- coding: utf-8 -*-
"""hsb(8).pdf 完整转换稿的最小内容精修（v3，只改有据可依之处）。

原则：格式一律不动（字体/字号/行距/页边距/表格/图片尺寸保持转换稿原样）；
只做内容层面修复：
1) 摘要按原文骨架轻度精简（保留全部关键数字与方法专名，压到第一页）；
2) 删除 pdf2docx 带入的 93 个纯数字页码碎片段（转换垃圾，非正文）；
3) 9.4 节复制 9.3 的笔误：替换为真正的问题四检验结论（内容取自 9.5/10.2 已有结论）；
4) 10.1 标题“模型的缺点”->“模型的优点”（该节内容本就是优点）；
5) 表8/表10 标题重复编号；
6) 图15 图注误标问题三 -> 问题二；
7) 9.3 正文“图30 和表16”错误引用 -> “图16 和表12”；
8) 5.5 末尾与上文完全重复的口径句删除；
9) 9.2 跨页断行“0.9%1.0%1.2%”粘连修复；
10) 图15(a)(c) 图注细节修正；
11) 删除 pdf2docx 每页一个的强制分节符（转换产物，让内容正常流动；
    字体/字号/行距/边距/表格/图片尺寸均不变）。
"""
import re
import docx
from docx.oxml.ns import qn

SRC = r"D:\数模竞赛\华数杯2026\A题项目\华数杯2026_A题\写作素材\hsb8_full_convert.docx"
DST = r"D:\数模竞赛\华数杯2026\A题项目\华数杯2026_A题\写作素材\华数杯论文_精修版.docx"

d = docx.Document(SRC)
paras = d.paragraphs
log = []


def set_para_text(par, text):
    if not par.runs:
        par.add_run(text)
        return
    par.runs[0].text = text
    for r in par.runs[1:]:
        r.text = ""


def find_par(prefix):
    for p in paras:
        if p.text.strip().startswith(prefix):
            return p
    return None


def remove_par(par):
    parent = par._p.getparent()
    if parent is not None:
        parent.remove(par._p)


# ---------- 1) 摘要轻度精简（原文骨架+全部关键数字） ----------
abstract_new = [
    "可靠且可复现的导通性评估对导电填充材料的设计与成本优化至关重要。本文针对微构体中"
    "随机填充圆柱介质A与球形介质B的导通判定、导通概率估计、临界体积分数反求与双介质"
    "最低成本设计问题，以三维最短距离＋并查集构建四问共用判定内核，结合蒙特卡洛、渗透"
    "模型拟合、概率二分、等效体积线性规划与Logistic代理优化求解，并将可行性结论统一置于"
    "95%置信区间下界不低于0.90的判据之下。",
    "针对问题一，将各行介质A抽象为轴线段，以左右带电面为源汇节点构造接触图，按题面判据"
    "换算轴线距离阈值，用向量化三维线段最近点算法生成接触边，以并查集判定连通并回溯跨接"
    "路径、定位断口。结果表明，主口径下组1因约279nm断口不导通，组2、组3均导通；同根"
    "合并口径三组全导通，作为敏感性对照列入第九章。",
    "针对问题二，在C1口径（介质完全位于盒内）下以拒绝采样生成随机构型，复用问题一内核"
    "做蒙特卡洛估计（每点3000次、多进程并行），以Wilson得分区间、收敛曲线与多种子复核"
    "刻画统计可靠性。结果表明，导通概率呈典型渗流相变，由0.5%时的0.060经1.0%时的0.730"
    "升至1.5%时的0.949；卷回口径（A1/B1）恒接近1，与题目非平凡设定不符，故主口径取C1。",
    "针对问题三，以Wilson区间下界不低于0.90为统一可行标准：先拟合Sigmoid与幂律截断渗透"
    "模型（加权最小二乘、AIC选优）并解析反解，再以整数介质数量为变量做概率二分，随后在"
    "两解邻域细化并多种子终验。最终定稿φ*=1.359%（N=961），终验P̂=0.911，95%CI="
    "[0.905,0.917]。",
    "针对问题四，构建“理论层＋工程层＋验证层”框架：理论层由纯A、纯B临界体积分数标定"
    "等效权重并建立线性规划求成本下界；工程层拟合Logistic代理曲面并搜索混合候选；验证层"
    "对全部候选做真实蒙特卡洛复核，低成本混合候选因P̂=0.818未达标被否决。最终最优方案"
    "为纯A，φ*=1.359%，成本14.27元；球形介质存在“面接触瓶颈”，纯B在60%体积分数下"
    "仍不满足判据。",
    "最后，对口径选择、导通阈值、样本量收敛、多种子复现及单价与等效权重扰动进行系统敏感"
    "性分析，全部结论稳健且可复现；方法可推广至复合材料导电网络、多孔介质渗流及电池电极"
    "导电剂网络等随机微结构体系。",
]
start = find_par("可靠且可复现的导通性评估")
kw = find_par("关键词")
if start is not None and kw is not None:
    pars = []
    on = False
    for p in paras:
        if p is start:
            on = True
        if on:
            pars.append(p)
        if p is kw:
            break
    for i, p in enumerate(pars):
        if i < len(abstract_new):
            set_para_text(p, abstract_new[i])
        else:
            remove_par(p)
    log.append(f"摘要：原文 {len(pars)} 段 -> 精简 {len(abstract_new)} 段（数字/方法专名全保留）")
else:
    log.append("摘要：未找到锚点！")


# ---------- 2) 删除纯数字页码碎片 ----------
removed = 0
for p in list(paras):
    if re.fullmatch(r"\d{1,3}", p.text.strip()):
        remove_par(p)
        removed += 1
log.append(f"页码碎片：删除 {removed} 段")


# ---------- 3) 9.4 复制笔误修复 ----------
h94 = None
for p in paras:
    if p.text.strip().startswith("9.4"):
        h94 = p
        break
if h94 is not None:
    runs = h94.runs
    if len(runs) >= 7:
        runs[3].text = "问题四模型的检验"
        new_body = (
            "问题四重点检验单价、等效体积权重扰动及纯B外推的不确定性。对A、B单价分别"
            "做±10%扰动并重解理论层线性规划，全部组合下LP极点均退化为纯A（成本"
            "12.84~15.70元），最优方案不随单价扰动改变；将等效体积权重w在±20%内扰动，"
            "LP最优恒为纯A（成本14.27元）。纯B扫描上限为60%，φ*_B≈61.4%为外推插值；"
            "50%、55%、60%三点真实MC的CI下界分别为0.721、0.817、0.856，均低于0.90，"
            "外推处P̂=0.902但CI下界0.882仍不达标，因此“纯B不达标”结论保守稳健，结果"
            "见图17。"
        )
        runs[6].text = new_body
        for r in runs[7:]:
            r.text = ""
    # 删除 9.4 与 9.5 之间的复制内容段
    content = []
    on = False
    for p in paras:
        if p is h94:
            on = True
            continue
        if on:
            if p.text.strip().startswith("9.5"):
                break
            content.append(p)
    n = 0
    for p in content:
        t = p.text.strip()
        if t.startswith("接触阈值缩放因子") or t.startswith("利用阈值敏感性曲线"):
            remove_par(p)
            n += 1
    log.append(f"9.4：正文替换为真问题四检验，删除复制段 {n} 段")
else:
    log.append("9.4：未找到！")


# ---------- 4) 10.1 标题 ----------
for p in paras:
    if p.text.strip().startswith("10.1"):
        for r in p.runs:
            if "模型的缺点" in r.text:
                r.text = r.text.replace("模型的缺点", "模型的优点", 1)
                log.append("10.1：标题“模型的缺点”->“模型的优点”")
                break


# ---------- 5) 表8/表10 ----------
for p in paras:
    t = p.text.strip()
    if t.startswith("表8:表8"):
        set_para_text(p, t.replace("表8:表8", "表8:", 1))
        log.append("表8：编号重复已修")
    if t.startswith("表10:表12"):
        set_para_text(p, t.replace("表10:表12", "表10:", 1))
        log.append("表10：编号重复已修")


# ---------- 6) 图15 图注 ----------
for p in paras:
    if p.text.strip().startswith("图15") and "问题三敏感性" in p.text:
        set_para_text(p, p.text.replace("问题三敏感性", "问题二敏感性", 1))
        log.append("图15：图注问题三->问题二")
        break


# ---------- 7) 图30/表16 引用 ----------
n = 0
for p in paras:
    t = p.text
    if "图30" in t or "表16" in t:
        t2 = t.replace("见图30 和表16", "见图16 和表12").replace("见图30和表16", "见图16和表12")
        if t2 != t:
            set_para_text(p, t2)
            n += 1
log.append(f"引用修正：{n} 处“图30/表16”->“图16/表12”")


# ---------- 8) 5.5 重复句 ----------
for p in paras:
    if p.text.strip().startswith("组1 的判定按附件行独立口径给出"):
        remove_par(p)
        log.append("5.5：删除与上文重复的口径句")
        break


# ---------- 9) 9.2 粘连/断行 ----------
for p in paras:
    if "其次，对φ=" in p.text:
        idx = paras.index(p)
        nxt = None
        for q in paras[idx + 1:]:
            if q.text.strip().startswith("0.9%1.0%1.2%"):
                nxt = q
                break
        if nxt is not None:
            tail = nxt.text.strip().replace("0.9%1.0%1.2%", "0.9%、1.0%、1.2%", 1)
            set_para_text(p, p.text + tail)
            remove_par(nxt)
            log.append("9.2：跨页断行已拼接，百分比顿号已补")
        break


# ---------- 10) 图15 图注细节 ----------
for p in paras:
    t = p.text
    if "M=500 3000" in t:
        set_para_text(p, t.replace("M=500 3000", "M=500~3000", 1))
        log.append("图15(a)：M=500~3000")
    if "（c）生成口径敏感性（A1/B1 卷回vs C1 完全悬\n浮）" in t:
        set_para_text(p, "（c）生成口径敏感性（A1/B1 卷回 vs C1 完全悬浮）")
        log.append("图15(c)：图注排版修正")


# ---------- 11) 删除 pdf2docx 强制分节符（仅页面镜像，非样式） ----------
body = d.element.body
nsect = 0
for p in body.iter(qn("w:p")):
    pPr = p.find(qn("w:pPr"))
    if pPr is not None:
        sectPr = pPr.find(qn("w:sectPr"))
        if sectPr is not None:
            pPr.remove(sectPr)
            nsect += 1
log.append(f"强制分节符：删除 {nsect} 个（页面镜像，非字体/行距/边距）")


d.save(DST)
print("\n".join(log))
print("SAVED:", DST)
