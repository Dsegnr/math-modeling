# -*- coding: utf-8 -*-
"""对转换后的完整论文 docx 做定点精修：
1) 摘要压缩到一页内；2) 修复 9.4 复制粘贴错误；3) 修复 10.1 标题；
4) 修复表8/图15 标题；5) 删 5.5 重复句；6) 图片宽度统一收缩。"""
import docx
from docx.shared import Cm

PATH = r"D:\数模竞赛\华数杯2026\A题项目\华数杯2026_A题\写作素材\hsb8_convert.docx"
d = docx.Document(PATH)


def set_para_text(par, text):
    """保留首 run 格式，整体替换段落文本。"""
    if not par.runs:
        par.add_run(text)
        return
    par.runs[0].text = text
    for r in par.runs[1:]:
        r.text = ""


def find_par(prefix):
    for p in d.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    return None


# ---------- 1) 摘要压缩（8 段替换为精简版） ----------
abstract_new = [
    "可靠且可复现的导通性评估对导电填充材料的设计与成本优化至关重要。本文针对微构体中"
    "随机填充圆柱介质A与球形介质B的导通判定、导通概率估计、临界体积分数反求与双介质"
    "最低成本设计问题，以三维最短距离与并查集构建四问共用的确定性判定内核，结合拒绝采样"
    "蒙特卡洛、渗透模型加权拟合、基于Wilson区间下界的概率二分、等效体积线性规划与Logistic"
    "代理优化求解，并将全部可行性结论统一置于“95%置信区间下界不低于0.90”的判据之下。",
    "针对问题一，将附件各行介质A抽象为轴线段，以左右带电面为源汇节点构造接触图，按题面"
    "表面间距判据换算轴线距离阈值，用向量化三维线段最近点算法批量生成接触边，以并查集"
    "判定连通并回溯跨接路径、定位断口。结果表明，主口径下组1因约279nm的关键断口不导通，"
    "组2、组3均存在完整通路；同根合并口径三组全导通，作为口径敏感性对照列入第九章。",
    "针对问题二，在介质完全位于盒内的C1口径下以拒绝采样生成随机构型，复用问题一内核做"
    "蒙特卡洛估计（每点3000次，多进程并行），以Wilson得分区间、收敛曲线与多种子复核刻画"
    "统计可靠性。结果表明，导通概率呈典型渗流相变，由0.5%时的0.060经1.0%时的0.730升至"
    "1.5%时的0.949；卷回口径（A1/B1）在相同区间恒接近1，与题目非平凡概率设定不符，故"
    "主口径取C1。",
    "针对问题三，以Wilson区间下界不低于0.90为统一可行标准：先拟合Sigmoid与幂律截断渗透"
    "模型（加权最小二乘、AIC选优）并解析反解，再以整数介质数量为变量做概率二分，随后在"
    "两解邻域高样本细化并多种子终验。最终定稿最低体积分数φ*=1.359%（N=961），终验"
    "P̂=0.911，95%CI=[0.905,0.917]，拟合质量另以残差诊断核对。",
    "针对问题四，构建“理论层＋工程层＋验证层”三层优化框架：理论层由纯A、纯B临界体积"
    "分数标定等效体积权重并建立线性规划求成本下界；工程层拟合Logistic代理曲面并以罚函数"
    "SLSQP搜索混合候选；验证层对全部候选做真实蒙特卡洛复核，低成本混合候选因此被否决。"
    "最终最优方案为纯A，φ*=1.359%，成本14.27元，不足纯B方案的一半；机理上球形介质存在"
    "“面接触瓶颈”，球体有效贴面厚度仅约1.8nm，纯B在60%体积分数下仍不满足判据。",
    "最后，对口径选择、导通阈值、样本量收敛、多种子复现及单价与等效权重扰动进行系统敏感性"
    "分析，全部结论稳健且可复现（固定随机流、逐样本落盘、断点续跑）；方法可推广至复合材料"
    "导电网络、多孔介质渗流及电池电极导电剂网络等随机微结构体系。",
]
start = find_par("可靠且可复现的导通性评估")
kw = find_par("关键词")
if start is not None and kw is not None:
    pars = []
    on = False
    for p in d.paragraphs:
        if p is start:
            on = True
        if on:
            pars.append(p)
        if p is kw:
            break
    for i, p in enumerate(pars):
        if i < len(abstract_new):
            set_para_text(p, abstract_new[i])
        else:
            p._p.getparent().remove(p._p)
    print("[abstract] replaced", len(pars), "paras")
else:
    print("[abstract] NOT FOUND", start is not None, kw is not None)


# ---------- 2) 修复 9.4（问题四检验）复制粘贴错误 ----------
h94 = find_par("9.4 问题四模型的检验")
if h94 is not None:
    new_text = (
        "问题四重点检验介质单价、等效体积权重扰动及纯B外推不确定性。对介质A、B单价分别"
        "做±10%扰动并重解理论层线性规划，全部组合下LP极点均退化为纯A（成本12.84~15.70元），"
        "最优方案不随单价扰动改变；将等效体积权重w在±20%内扰动，LP最优恒为纯A（成本14.27元），"
        "说明成本结论对w不敏感。纯B扫描上限为60%，φ*_B≈61.4%为外推插值；50%、55%、60%三点"
        "真实MC结果的CI下界分别为0.721、0.817、0.856，均低于0.90，外推处P̂=0.902但CI下界"
        "0.882仍不达标，因此“纯B不达标”结论保守稳健。相关结果见图17与表15（图号以正文实际"
        "编号为准）。"
    )
    content = []
    on = False
    for p in d.paragraphs:
        if p is h94:
            on = True
            continue
        if on:
            if p.text.strip().startswith("9.5"):
                break
            content.append(p)
    if content:
        set_para_text(content[0], new_text)
        for p in content[1:]:
            p._p.getparent().remove(p._p)
        print("[9.4] fixed, removed", len(content) - 1, "paras")
else:
    print("[9.4] NOT FOUND")


# ---------- 3) 10.1 标题：模型的缺点 -> 模型的优点 ----------
h101 = find_par("10.1 模型的缺点")
if h101 is not None:
    set_para_text(h101, h101.text.replace("10.1 模型的缺点", "10.1 模型的优点", 1))
    print("[10.1] heading fixed")


# ---------- 4) 表8 表题重复 ----------
for p in d.paragraphs:
    if "表8" in p.text and "三方法" in p.text and p.text.strip().startswith("表8 表8"):
        set_para_text(p, p.text.replace("表8 表8", "表8", 1))
        print("[表8] caption fixed")
        break


# ---------- 5) 图15 图注：问题三 -> 问题二 ----------
for p in d.paragraphs:
    if "问题三敏感性分析结果" in p.text:
        set_para_text(p, p.text.replace("问题三敏感性分析结果", "问题二敏感性分析结果", 1))
        print("[图15] caption fixed")
        break


# ---------- 6) 5.5 重复口径句精简 ----------
for p in d.paragraphs:
    t = p.text
    if t.strip().startswith("合并口径（同根圆柱先合并为一个节点）下三组全部导通"):
        cut = "组1 的判定按附件行独立口径给出；若按同根合并的物理口径，组1 亦导通（三组全导通），该口径差异已在第九章敏感性中对照分析。"
        if cut in t:
            set_para_text(p, t.replace(cut, "", 1).rstrip("。") + "。")
            print("[5.5] duplicate sentences removed")
        break


# ---------- 7) 图片宽度收缩（>13cm 等比缩小） ----------
from docx.shared import Emu
target = Cm(13.0)
shrunk = 0
for sh in d.inline_shapes:
    if sh.width and sh.width > target:
        ratio = target / sh.width
        sh.width = target
        sh.height = Emu(int(sh.height * ratio))
        shrunk += 1
print("[images] shrunk", shrunk)

d.save(PATH)
print("SAVED")
