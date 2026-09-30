# -*- coding: utf-8 -*-
"""修正公式编号连续性：九章 29~36 续排；五~八章正文交叉引用同步。"""
import re
import docx

MATH = "{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath"


def renumber_offset(path, offset):
    doc = docx.Document(path)
    n = 0
    for p in doc.paragraphs:
        if p._p.find(MATH) is not None:
            n += 1
            for r in p.runs:
                if re.search(r"（公式\s*\d+）", r.text):
                    r.text = re.sub(r"（公式\s*\d+）",
                                    "（公式%d）" % (offset + n), r.text)
    doc.save(path)
    return n


def fix_crossref(path, old, new):
    doc = docx.Document(path)
    hit = 0
    for p in doc.paragraphs:
        for r in p.runs:
            if old in r.text and r._element.find(MATH) is None:
                r.text = r.text.replace(old, new)
                hit += 1
    doc.save(path)
    return hit


if __name__ == "__main__":
    p9 = (r"D:\数模竞赛\华数杯2026\A题项目\华数杯2026_A题\写作素材"
          r"\第九章_模型的分析与检验_初稿.docx")
    p58 = (r"D:\数模竞赛\华数杯2026\A题项目\华数杯2026_A题\写作素材"
           r"\五至八章_模型建立与求解_初稿.docx")
    n9 = renumber_offset(p9, 28)
    h = fix_crossref(p58, "按公式10取最小可行 N", "按公式18取最小可行 N")
    print("ch9 renumbered to", n9, "formulas (29..%d)" % (28 + n9))
    print("crossref fixed:", h)
